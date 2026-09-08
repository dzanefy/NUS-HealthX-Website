from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='/Users/dzanefy/Desktop/HealthX/NUS HealthX Website Mock/HealthX Accelerate Team Meeting Worksheet.docx'
navy='0B3558'; teal='008A8A'; gray='D9D9D9'
def shade(c,fill):
 p=c._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),fill); p.append(x)
def borders(t):
 p=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
 for e in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),gray); b.append(x)
 p.append(b)
def put(c,txt,b=False,col=None):
 c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(txt); r.font.name='Aptos'; r.font.size=Pt(9); r.bold=b
 if col:r.font.color.rgb=RGBColor.from_string(col)
 c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(d,heads,rows):
 t=d.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
 for i,h in enumerate(heads): put(t.rows[0].cells[i],h,True,'FFFFFF'); shade(t.rows[0].cells[i],navy)
 for j,row in enumerate(rows):
  cs=t.add_row().cells
  for i,v in enumerate(row): put(cs[i],v)
  if j%2:
   for c in cs: shade(c,'F5F8FA')
 d.add_paragraph().paragraph_format.space_after=Pt(1)
def head(d,text,n=1):
 p=d.add_paragraph(style='Heading '+str(n)); p.paragraph_format.space_before=Pt(11 if n==1 else 7); p.paragraph_format.space_after=Pt(4); p.add_run(text)
def para(d,text):
 p=d.add_paragraph(text); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.05
def bullets(d,items):
 for s in items:
  p=d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(s)

d=Document(); s=d.sections[0]; s.top_margin=Inches(.6); s.bottom_margin=Inches(.55); s.left_margin=Inches(.65); s.right_margin=Inches(.65)
d.styles['Normal'].font.name='Aptos'; d.styles['Normal'].font.size=Pt(10); d.styles['Normal'].font.color.rgb=RGBColor.from_string('263746')
for n,z,c in [('Title',25,navy),('Heading 1',16,navy),('Heading 2',12,teal)]:
 st=d.styles[n]; st.font.name='Aptos Display'; st.font.size=Pt(z); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c)
d.add_paragraph('HealthX Accelerate Team Meeting Worksheet',style='Title')
x=d.add_paragraph('Questions, feasibility review, and implementation planning'); x.runs[0].italic=True; x.runs[0].font.size=Pt(13); x.runs[0].font.color.rgb=RGBColor.from_string(teal)
para(d,'Purpose: use this worksheet to understand the X’ccelerate team’s draft, confirm what they need, identify risks early, and leave the meeting with agreed next steps. Fill in the answer and decision columns during or after the meeting.')
head(d,'Meeting details'); table(d,['Item','Notes'],[['Date and time',''],['Attendees and roles',''],['Meeting owner',''],['Target pilot or launch date',''],['Documents or links received','']])
head(d,'1 Executive view of the proposed implementation')
para(d,'The draft proposes one HealthX website with three connected user groups: students, companies, and HealthX administrators. Students create profiles and discover opportunities; companies create profiles and submit internships; HealthX verifies, publishes, matches, shortlists, and tracks placements.')
table(d,['Area','What the draft proposes','Feasibility','Confirm'],[['Student access','NUS login, profile, resume, skills, interests, browse, indicate interest','High','Login, fields, privacy, resume access'],['Company access','Company login, profile, internship posting, review submission','Medium to high','Registration, verification, role review'],['Listings','Internship listings and company directory','High','Approval, filters, expiry, duplicates'],['Matching','Skills, interests, availability, requirements','Medium','Rules, data quality, human override'],['Admin','Review applications and placement stages','High','Queues, permissions, statuses']])
head(d,'2 Questions to ask before agreeing to scope')
table(d,['Question','Why it matters','Answer / notes'],[['Who is the first pilot group and how many users?','Sets size, support, and realistic scope.',''],['What is the one outcome the platform must achieve?','Prevents feature creep.',''],['Which features are mandatory for launch?','Separates MVP from later work.',''],['Will staff still use email or spreadsheets?','Reveals hidden workflow needs.',''],['What is the expected volume of roles and applications?','Affects design and infrastructure.',''],['Who approves companies, roles, matches, and offers?','Clarifies accountability.',''],['What is the support process for user problems?','A live service needs an owner.',''],['What does success look like after the pilot?','Creates acceptance criteria.','']])
head(d,'3 Module implementation discussion')
table(d,['Module','Suggested implementation flow','Questions / requirements to confirm','Notes'],[['1 Student profile','Login → onboarding → resume → skills/interests → availability → preview','Required fields, NUS identity, consent, file limits, visibility',''],['2 Company directory','Company form → verification queue → approved profile → directory','Evidence, reviewer, public fields, re-verification',''],['3 Internship board','Draft role → requirements → submit → review → publish → expire','Required fields, closing date, edits after publishing',''],['4 Matching','Tags → compare profile and role → hard filters → rank → explain → review','Weights, missing data, bias checks, manual override',''],['5 Admin','Review queues → shortlist → interviews → offers → placement → reports','Statuses, notifications, audit log, ownership','']])
head(d,'4 Pros, cons, and feasibility')
table(d,['Aspect','Advantages','Risks or drawbacks','Conclusion'],[['Single shared website','One place for all users; easy to explain.','Permission and workflow complexity.',''],['NUS login','Reduces fake accounts.','May require NUS approval and SSO work.',''],['Verified directory','Builds trust.','Needs ongoing staff maintenance.',''],['Resume and profile data','Supports matching and company review.','Sensitive data needs strict controls.',''],['Matching engine','Reduces manual searching.','Bad data may create poor or biased results.',''],['Admin workflow','Quality control and auditability.','Review workload may become a bottleneck.',''],['Placement tracking','Makes outcomes visible.','Requires consistent updates.','']])
head(d,'5 Recommended feasibility position')
para(d,'Overall, the concept is feasible as a phased web platform. Profiles, directories, listings, and admin workflows are conventional web features. The main feasibility questions are NUS access, privacy approval, company verification, operational ownership, and matching data quality.')
table(d,['Recommendation','Reason','Decision'],[['Proceed with a small MVP pilot','Core workflow can be tested without advanced automation.',''],['Treat matching as decision support in version one','Keeps results explainable and reviewable.',''],['Confirm privacy and authentication first','These choices affect the technical design.',''],['Make admin workflow a first-class feature','The process fails if reviews have no owner.',''],['Deliver in phases','Reduces risk and allows feedback.','']])
head(d,'6 Technical and operational questions')
table(d,['Question','Answer / owner / due date'],[['Which authentication method is approved for NUS students?',''],['Where may profiles and resumes be stored?',''],['Who can view, download, edit, or delete resumes?',''],['How long are inactive records retained?',''],['How are external companies verified?',''],['Are email, calendar, or NUS integrations required?',''],['Who maintains skills and role categories?',''],['Who handles incidents, complaints, and recovery?',''],['What reports are required?',''],['What budget, people, and timeline are available?','']])
head(d,'7 Decisions to leave the meeting with')
table(d,['Decision','Agreed answer','Owner','Due date'],[['MVP features included','','',''],['Pilot size and participants','','',''],['Authentication approach','','',''],['Data and resume storage','','',''],['Company verification process','','',''],['Matching rules version one','','',''],['Admin statuses and owners','','',''],['Design review date','','',''],['Pilot date','','','']])
head(d,'8 Action items and follow-up')
table(d,['Action','Owner','Due date','Status'],[['Send final requirements or examples','','',''],['Confirm privacy and login requirements','','',''],['Provide sample data','','',''],['Review first prototype','','',''],['Agree pilot testers','','',''],['Schedule next meeting','','','']])
head(d,'Appendix A Suggested first release'); bullets(d,['Student access, profile, resume, skills, interests, and availability.','Company profile, verification, internship creation, review, and publishing.','Opportunity browsing, role details, indicate interest, shortlist, and status tracking.','Admin dashboard for queues, matching suggestions, and placement progress.','Basic confirmations and support instructions.'])
head(d,'Appendix B Later enhancements'); bullets(d,['Advanced AI matching and resume parsing.','Messaging, calendar scheduling, contracts, payments, and automated offers.','Mobile app, employer ratings, advanced analytics, and external job integrations.'])
d.core_properties.title='HealthX Accelerate Team Meeting Worksheet'; d.core_properties.subject='Questions, feasibility review, and implementation planning'; d.save(OUT); print(OUT)
