from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT = "/Users/dzanefy/Desktop/HealthX/NUS HealthX Website Mock/HealthX Accelerate Implementation Guide.docx"
navy = "0B3558"
teal = "008A8A"
light_blue = "EAF4F8"
light_teal = "E9F6F3"
gray = "D9D9D9"

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def borders(table):
    tblPr = table._tbl.tblPr; b = tblPr.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge; el = b.find(qn(tag))
        if el is None: el = OxmlElement(tag); b.append(el)
        el.set(qn('w:val'),'single'); el.set(qn('w:sz'),'4'); el.set(qn('w:color'),gray)
def set_cell(cell, text, bold=False, color=None):
    cell.text = ''; p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    r = p.add_run(text); r.bold = bold; r.font.size = Pt(9.5); r.font.name = 'Aptos'
    if color: r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(doc, headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers)); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = True; borders(t)
    for i,h in enumerate(headers): set_cell(t.rows[0].cells[i], h, True, 'FFFFFF'); shade(t.rows[0].cells[i], navy)
    for ri,row in enumerate(rows):
        cells=t.add_row().cells
        for i,val in enumerate(row): set_cell(cells[i], val)
        if ri%2==1:
            for c in cells: shade(c, 'F5F8FA')
    if widths:
        for row in t.rows:
            for i,w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t
def heading(doc, text, level=1):
    p=doc.add_paragraph(style=f'Heading {level}'); p.paragraph_format.space_before=Pt(12 if level==1 else 7); p.paragraph_format.space_after=Pt(5)
    p.add_run(text); return p
def para(doc, text, boldlead=None):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(6); p.paragraph_format.line_spacing=1.08
    if boldlead and text.startswith(boldlead): p.add_run(boldlead).bold=True; p.add_run(text[len(boldlead):])
    else: p.add_run(text)
    return p
def bullets(doc, items):
    for x in items:
        p=doc.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(x)

doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.65); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.75); sec.right_margin=Inches(.75)
styles=doc.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].font.color.rgb=RGBColor.from_string('263746')
for name,size,color in [('Title',28,navy),('Heading 1',17,navy),('Heading 2',12,teal)]:
    s=styles[name]; s.font.name='Aptos Display'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor.from_string(color)

p=doc.add_paragraph(style='Title'); p.add_run('HealthX Accelerate Implementation Guide')
p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(16); r=p.add_run('A simple starting plan for building the MedTech internship platform'); r.italic=True; r.font.size=Pt(14); r.font.color.rgb=RGBColor.from_string(teal)
para(doc, 'Purpose: turn the HealthX Accelerate concept into a practical first release that connects NUS students with verified MedTech companies and internship opportunities. This guide explains what to decide before development, what to deliver in each phase, and what “ready to start” means.')
heading(doc,'The recommended starting point',1)
para(doc,'Start with a small pilot, not the full platform. The first release should let a student create a profile and show interest in an opportunity; let a company create a profile and submit an internship; and let HealthX review, match, and manage the process. Build the five core modules around this single journey, then improve it using pilot feedback.')
table(doc,['First release must do','Can wait until later'],[
['NUS student sign-in or approved student access','Advanced recommendation AI'],['Student profile, resume, skills, interests, availability','In-platform chat and complex notifications'],['Company profile and verification status','Automated offer and placement contracts'],['Internship posting, review, and publishing','Analytics dashboards and integrations'],['Browse, indicate interest, shortlist, and track status','Mobile app']],[2.8,3.7])

heading(doc,'1 What the diagram is asking the platform to do',1)
para(doc,'The model has three user groups and one shared workflow. Students discover and apply; companies describe opportunities and review candidates; HealthX administrators verify, publish, match, and manage the process. The website is the shared entry point.')
table(doc,['User','Main journey','Success result'],[
['Student','Log in → build profile → upload resume → add skills and interests → browse roles → indicate interest','A complete, useful student profile and a clear expression of interest'],
['Company','Log in → create company profile → post role → add requirements → submit for review','A verified organisation with a clear, publishable opportunity'],
['HealthX admin','Review companies and roles → publish → support matching → shortlist → track interviews, offers, and placements','A trusted, auditable placement process']],[1.15,3.65,1.7])

heading(doc,'2 Decisions to make before building',1)
para(doc,'These decisions prevent rework. Write the answers down and get agreement from the HealthX owner, student representative, and company or programme representative before development begins.')
table(doc,['Decision area','Questions to answer','Deliverable'],[
['Scope','What is included in the pilot? Which items are explicitly out of scope?','One-page MVP scope'],['Users and access','Who may register? How will NUS identity be verified? Who can approve a company?','Roles and access rules'],['Process','What does “interested,” “shortlisted,” “interview,” “offer,” and “placed” mean?','Status and workflow map'],['Data and privacy','Which student data is required? Who can see resumes and contact details? How long is data kept?','Data and privacy rules'],['Matching','Which skills, interests, availability, and requirements matter? Is matching advice only or a decision?','Matching rules v1'],['Governance','Who reviews content, handles disputes, and owns the system after launch?','Owner and operating model']],[1.1,3.55,1.85])

heading(doc,'3 Deliverables needed before development starts',1)
para(doc,'The project is ready to start when the following items exist, have an owner, and are approved. “Draft” is acceptable for some items, but unknowns should be recorded rather than left implicit.')
table(doc,['Deliverable','What it contains','Owner'],[
['MVP scope','Pilot audience, included features, exclusions, target launch date','Programme lead'],['User journeys','Step-by-step journeys for student, company, and admin','Product lead'],['Requirements list','What each screen and action must do; priority: must, should, later','Product + engineering'],['Data dictionary','Fields for profiles, resumes, companies, roles, interests, and statuses','Product + data/privacy'],['Wireframes','Low-fidelity screens for the complete happy path and key error states','UX/UI'],['Content pack','Field labels, help text, review messages, emails, and empty states','Content owner'],['Access and privacy plan','Login method, permissions, consent, retention, and incident process','NUS/HealthX owner'],['Matching definition','Inputs, scoring or rules, explanation shown to users, and manual override','Product + programme'],['Pilot plan','Participants, test tasks, feedback method, support contact, and success measures','Programme lead'],['Technical approach','Architecture, hosting, file storage, database, audit logs, and environments','Technical lead']],[1.35,3.9,1.25])

heading(doc,'4 Build the platform in five modules',1)
modules=[('Module 1 Student profile','Create profile fields, resume upload, skills, interests, availability, consent, and profile completeness.','Student can save, edit, preview, and control what is shared.'),('Module 2 Company directory','Create company records, verification status, organisation details, sector, website, and contact person.','Only approved companies are visible as verified organisations.'),('Module 3 Internship board','Create draft roles, requirements, dates, location, commitment, application instructions, and review status.','Admin can review and publish a clear opportunity.'),('Module 4 Matching','Use a transparent first version based on skills, interests, availability, and role requirements. Allow admin review and override.','Users see relevant suggestions with a simple explanation of why they match.'),('Module 5 Accelerate admin','Provide queues for company review, role review, student interest, shortlist, interview, offer, and placement.','HealthX can manage the lifecycle without spreadsheets becoming the source of truth.')]
for title,build,done in modules:
    heading(doc,title,2); para(doc,build,'Build: '); para(doc,done,'Done when: ')

heading(doc,'5 Suggested delivery phases',1)
table(doc,['Phase','Work','Output'],[
['0 Align','Confirm scope, owners, rules, data, and pilot','Approved project brief'],['1 Design','Map journeys, create wireframes, write content, review privacy','Clickable prototype and requirements'],['2 Build foundation','Set up access, data model, profiles, file handling, audit trail','Working student and company foundations'],['3 Build workflow','Internship board, admin review, matching v1, interest and status tracking','End-to-end pilot workflow'],['4 Test and pilot','Run usability, permission, security, and data tests; pilot with a small cohort','Pilot report and prioritised fixes'],['5 Launch and improve','Train admins, publish support process, monitor outcomes, iterate','Live service and improvement backlog']],[1.1,3.6,1.8])

heading(doc,'6 Minimum screens for the first release',1)
bullets(doc,['Landing page with separate student and company entry points','Student sign-in, onboarding, profile editor, resume upload, and profile preview','Opportunity list, filters, opportunity detail, and indicate-interest action','Company sign-in, onboarding, company profile, role editor, preview, and submit-for-review action','Admin dashboard with review queues, detail views, match suggestions, shortlist, and status updates','Basic confirmation, error, empty, consent, and support states'])

heading(doc,'7 Matching version one',1)
para(doc,'Keep the first matching approach understandable. Convert each opportunity requirement and each student skill or interest into a shared list of tags. Compare the lists, then adjust for availability and any hard requirements. The result should be a recommendation for review, not an automatic rejection or final decision.')
table(doc,['Input','Example','First-release treatment'],[
['Skills','CAD, Python, clinical research','Strong match when tags overlap'],['Interests','Diagnostics, digital health','Adds relevance to the ranking'],['Availability','June to August, part-time','Filter or reduce ranking when incompatible'],['Role requirements','Must be NUS student; 10 weeks','Hard checks before recommendation'],['Human review','Admin judgement and company feedback','Can override or correct the suggestion']],[1.4,2.1,3.0])

heading(doc,'8 Testing checklist before launch',1)
bullets(doc,['A student can complete the journey using only the interface and understands what happens next.','A company cannot publish before required information and review are complete.','An unverified company or unpublished role is not shown to students.','Permissions prevent companies from seeing private student data or other companies’ records.','Resume upload rejects unsafe or unsupported files and does not expose the storage link publicly.','Every status change has an owner, timestamp, and clear next action.','Matching suggestions can be explained in plain language and manually corrected.','Admin can export or recover key records and has a support procedure for mistakes.'])

heading(doc,'9 Pilot success measures',1)
para(doc,'Measure the pilot around completion and usefulness, not just the number of screens built.')
table(doc,['Measure','How to check','Good signal'],[
['Student completion','% who finish a usable profile','Most pilot students reach profile complete'],['Company readiness','% of submitted roles approved without major rework','Roles are clear and complete'],['Time to publish','Days from company submission to live role','Review does not depend on manual chasing'],['Match usefulness','Student and company rating after review','Suggestions are relevant and explainable'],['Workflow visibility','% of active cases with a current status','No case is lost between interest and placement'],['Trust and safety','Permission, privacy, and upload test results','No critical access or data-handling issues']],[1.5,3.3,1.7])

heading(doc,'10 Immediate next steps',1)
for i,x in enumerate(['Name one HealthX product owner and one technical owner.','Agree the pilot audience and choose a small group of students and companies.','Hold a 60-minute scope workshop using the decisions in Section 2.','Create the MVP scope, user journeys, field list, and status map.','Sketch and review the minimum screens with three students and two companies.','Only then confirm the build estimate and start implementation.'],1):
    p=doc.add_paragraph(style='List Number'); p.paragraph_format.space_after=Pt(3); p.add_run(x)
heading(doc,'11 Software and tools you will need',1)
para(doc,'You do not need to buy or learn every tool at once. The recommended setup below is enough to design, build, test, and run the first pilot. Start with the tools marked Needed now; add the others only when the project reaches that phase.')
table(doc,['Tool','What you use it for','Priority'],[
['Figma','Design wireframes, clickable prototypes, user flows, and the visual interface.','Needed now'],
['React and Vite','Build the website interface and connect the main screens. The current HealthX project already uses this setup.','Needed for build'],
['Tailwind CSS','Style the website consistently and quickly using reusable design rules.','Needed for build'],
['Node.js and npm or pnpm','Run the development project, install packages, and create production builds.','Needed for build'],
['Git and GitHub','Save versions of the code, review changes, and collaborate safely.','Needed for build'],
['A code editor such as VS Code','Write and organise the website code.','Needed for build'],
['Database such as Supabase','Store student profiles, companies, internships, interests, statuses, and permissions.','Needed for pilot'],
['File storage such as Supabase Storage','Store resumes securely with access controls.','Needed for pilot'],
['Authentication such as NUS SSO or Supabase Auth','Allow students, companies, and admins to sign in with the correct permissions.','Needed for pilot'],
['Email service such as Resend or SendGrid','Send confirmations, review updates, interview messages, and placement notifications.','Useful for pilot'],
['Analytics such as PostHog or Google Analytics','Understand where users drop off and which features are used.','Useful later'],
['Issue tracking such as GitHub Issues, Linear, or Trello','Track bugs, decisions, tasks, and the delivery backlog.','Needed now'],
['Testing tools such as Playwright','Check that sign-in, uploads, permissions, matching, and status changes keep working.','Needed before launch'],
['Password manager and shared documentation such as Notion or Google Drive','Store decisions, process instructions, credentials, and training materials safely.','Needed now']],[1.65,4.1,1.0])
heading(doc,'A simple recommended setup',2)
bullets(doc,['Plan and design: Figma plus Notion or Google Drive.','Build: VS Code, GitHub, Node.js, React, Vite, and Tailwind CSS.','Run the pilot: Supabase for database, authentication, and resume storage; an email service for notifications.','Test and improve: Playwright for repeatable checks, GitHub Issues for bugs, and a simple analytics tool for usage.'])
para(doc,'Important: confirm NUS requirements before selecting authentication, hosting, email, analytics, or resume storage. Student information and resumes should only be stored in services approved for the project, with clear access permissions and retention rules.')
heading(doc,'Appendix A Definition of ready',1)
para(doc,'Development should begin when the team can explain the complete student, company, and admin journeys; has agreed the required data and permissions; has approved the pilot scope; and has someone responsible for every review and support step. If any of these is missing, resolve it first or record it as an explicit pilot risk.')
heading(doc,'Appendix B Plain language glossary',1)
table(doc,['Term','Meaning'],[['MVP','The smallest useful first version of the platform.'],['Pilot','A limited trial with selected users before wider launch.'],['Role','An internship or project opportunity posted by a company.'],['Match','A recommendation that compares student information with role needs.'],['Admin','A HealthX team member who reviews and manages the platform.'],['Status','The current step in a process, such as draft, reviewed, shortlisted, or placed.']],[1.2,5.3])

doc.core_properties.title='HealthX Accelerate Implementation Guide'; doc.core_properties.subject='Simple implementation plan and pre-development deliverables'; doc.core_properties.author='HealthX'
doc.save(OUT)
print(OUT)
