from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='/Users/dzanefy/Desktop/HealthX/NUS HealthX Website Mock/Improved Website Pillar Requirements.docx'
navy='0B3558'; teal='008A8A'; gray='D9D9D9'; pale='F5F8FA'
def shade(c,fill):
 p=c._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),fill); p.append(x)
def borders(t):
 p=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
 for e in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),gray); b.append(x)
 p.append(b)
def put(c,text,b=False,col=None):
 c.text=''; p=c.paragraphs[0]; p.paragraph_format.space_after=Pt(2); r=p.add_run(text); r.font.name='Aptos'; r.font.size=Pt(9.5); r.bold=b
 if col:r.font.color.rgb=RGBColor.from_string(col)
 c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
def table(d,heads,rows):
 t=d.add_table(rows=1,cols=len(heads)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; borders(t)
 for i,h in enumerate(heads): put(t.rows[0].cells[i],h,True,'FFFFFF'); shade(t.rows[0].cells[i],navy)
 for j,row in enumerate(rows):
  cs=t.add_row().cells
  for i,v in enumerate(row): put(cs[i],v)
  if j%2:
   for c in cs: shade(c,pale)
 d.add_paragraph().paragraph_format.space_after=Pt(1)
def head(d,text,n=1):
 p=d.add_paragraph(style='Heading '+str(n)); p.paragraph_format.space_before=Pt(11 if n==1 else 7); p.paragraph_format.space_after=Pt(4); p.add_run(text)
def para(d,text):
 p=d.add_paragraph(text); p.paragraph_format.space_after=Pt(5); p.paragraph_format.line_spacing=1.05
def bullets(d,items):
 for s in items:
  p=d.add_paragraph(style='List Bullet'); p.paragraph_format.space_after=Pt(2); p.add_run(s)

d=Document(); s=d.sections[0]; s.top_margin=Inches(.65); s.bottom_margin=Inches(.6); s.left_margin=Inches(.7); s.right_margin=Inches(.7)
d.styles['Normal'].font.name='Aptos'; d.styles['Normal'].font.size=Pt(10.5); d.styles['Normal'].font.color.rgb=RGBColor.from_string('263746')
for n,z,c in [('Title',25,navy),('Heading 1',16,navy),('Heading 2',12,teal)]:
 st=d.styles[n]; st.font.name='Aptos Display'; st.font.size=Pt(z); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c)
d.add_paragraph('Website Pillar Requirements Workbook',style='Title')
x=d.add_paragraph('A fill-in template for planning each HealthX pillar page'); x.runs[0].italic=True; x.runs[0].font.size=Pt(13); x.runs[0].font.color.rgb=RGBColor.from_string(teal)
para(d,'Please complete one section for your pillar. The Technology Team will use your answers to plan the page structure, write or edit content, identify required features, and decide what needs to be designed or built. You do not need to know the technical solution. Describe what you want users to see and do; we will translate it into the website design.')
head(d,'How to complete this workbook')
bullets(d,['Fill in the blank cells with short, specific answers. Add links to existing slides, documents, images, or examples where helpful.','If you have no preference, write “No preference” or “Technology Team to propose”.','For features, describe the user, the action, and the desired result. Example: “A visitor can register interest in an upcoming workshop and receive confirmation.”','Mark each request as Must have, Should have, or Nice to have. This helps us plan a feasible first release.','Return the completed section to the Technology Team before the page design review.'])
head(d,'Information to provide once for all pillars')
table(d,['Item','Response'],[['Person completing this workbook',''],['Pillar name or team',''],['Main contact for follow-up',''],['Target date for first draft',''],['Links to AGM slides or existing materials',''],['Any words, claims, images, or information that must not be changed','']])
head(d,'Shared guidance for all pillar pages')
para(d,'Every pillar page can include a clear introduction, key content, sub-pillars or focus areas, relevant people or partners, events, calls to action, and supporting media. The final layout will be based on the information you provide and the overall HealthX website design system so that all pillars feel like part of one site.')

for idx in range(1,6):
 d.add_page_break(); head(d,f'Pillar {idx} Requirements')
 para(d,'Pillar name: ________________________________________________')
 table(d,['Requirement area','What the pillar team should provide','Your response'],[
 ['Purpose and audience','What is this pillar about? Who should the page help: students, companies, researchers, partners, or the public?',''],
 ['Short description','Write a 1–3 sentence introduction for the top of the page. Mention the pillar’s role and any sub-pillars.',''],
 ['Key content','List the most important information, achievements, programmes, people, partners, or resources to feature.',''],
 ['Sub-pillars or sections','List any sub-pillars or themes. For each, give a one-sentence description and the order of importance.',''],
 ['People and partners','Which advisors, leads, researchers, companies, student teams, or external partners should be shown? What information may be published?',''],
 ['Page layout preference','Describe your preferred order of sections. Attach a sketch, AGM slide, or example website if useful. If you have no preference, write “Technology Team to propose”.',''],
 ['Visual direction','Preferred images, icons, colours, diagrams, video, or examples. Note any brand or accessibility requirements.',''],
 ['Call to action','What should a visitor do next? Examples: sign up, contact the team, view a programme, join an event, read a resource, or partner with HealthX.',''],
 ['Events for the shared timeline','Which events should appear on the general HealthX timeline? Include event name, date, time, location or link, audience, and whether it is public.',''],
 ['Other content needs','News, testimonials, FAQs, downloadable resources, forms, external links, or anything else not covered above.','']])
 head(d,'Features requested by this pillar',2)
 para(d,'Complete one row per feature. Features may be simple content elements or interactive functions. The Technology Team will assess feasibility, privacy, maintenance, and whether each feature belongs in the first release.')
 table(d,['Feature or function','Who uses it and what they do','What should happen','Priority','Notes or examples'],[
 ['','','','Must / Should / Nice',''],['','','','Must / Should / Nice',''],['','','','Must / Should / Nice',''],['','','','Must / Should / Nice','']])
 head(d,'Content and maintenance responsibilities',2)
 table(d,['Question','Response'],[['Who owns the content after launch?',''],['How often should it be reviewed?',''],['Who approves updates before publication?',''],['Are there time-sensitive items that need expiry dates?',''],['What information should visitors be able to search or filter?',''],['Any privacy, copyright, consent, or accessibility concerns?','']])
 head(d,'Technology Team review notes',2)
 table(d,['Item','Technology Team notes'],[['Proposed page structure',''],['Features feasible for first release',''],['Features recommended for later',''],['Information still required',''],['Follow-up owner and date','']])

d.core_properties.title='Website Pillar Requirements Workbook'; d.core_properties.subject='Fill-in requirements template for HealthX pillar pages'; d.save(OUT); print(OUT)
