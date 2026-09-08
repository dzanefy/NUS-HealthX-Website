from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='/Users/dzanefy/Desktop/HealthX/NUS HealthX Website Mock/Simple Website Pillar Needs Form.docx'
navy='0B3558'; teal='008A8A'; gray='D9D9D9'
def shade(c,f):
 p=c._tc.get_or_add_tcPr(); x=OxmlElement('w:shd'); x.set(qn('w:fill'),f); p.append(x)
def table(d,rows):
 t=d.add_table(rows=0,cols=2); t.alignment=WD_TABLE_ALIGNMENT.CENTER
 p=t._tbl.tblPr; b=OxmlElement('w:tblBorders')
 for e in ('top','left','bottom','right','insideH','insideV'):
  x=OxmlElement('w:'+e); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),gray); b.append(x)
 p.append(b)
 for i,(a,btxt) in enumerate(rows):
  cs=t.add_row().cells
  for c,txt in zip(cs,[a,btxt]):
   c.text=txt
   for r in c.paragraphs[0].runs: r.font.name='Aptos'; r.font.size=Pt(10)
  if i==0:
   for c in cs: shade(c,navy)
   for r in cs[0].paragraphs[0].runs+runs(cs[1]): r.font.color.rgb=RGBColor(255,255,255); r.bold=True
  else: shade(cs[0],'F5F8FA')
 d.add_paragraph()
def runs(c): return c.paragraphs[0].runs
def head(d,text):
 p=d.add_paragraph(style='Heading 1'); p.paragraph_format.space_before=Pt(12); p.add_run(text)

d=Document(); s=d.sections[0]; s.top_margin=Inches(.7); s.bottom_margin=Inches(.65); s.left_margin=Inches(.75); s.right_margin=Inches(.75)
d.styles['Normal'].font.name='Aptos'; d.styles['Normal'].font.size=Pt(11); d.styles['Normal'].font.color.rgb=RGBColor.from_string('263746')
for n,z,c in [('Title',25,navy),('Heading 1',16,navy)]:
 st=d.styles[n]; st.font.name='Aptos Display'; st.font.size=Pt(z); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c)
d.add_paragraph('Website Pillar Needs Form',style='Title')
p=d.add_paragraph('A simple form for collecting the initial ideas of each pillar team'); p.runs[0].italic=True; p.runs[0].font.color.rgb=RGBColor.from_string(teal)
d.add_paragraph('Hi pillar teams! Please use the table for your pillar to share the general content and features you would like on the HealthX website. You do not need to provide technical details or final wording yet. Short answers, bullet points, links, or example websites are welcome. If you have no preference, write “No preference”.')
d.add_paragraph('The Technology Team will use your responses to propose the page layout and decide how the requested features can be implemented.')
table(d,[('General details','Your response'),('Pillar name',''),('Person completing this form',''),('Best contact for follow-up',''),('Link to AGM slides or existing materials','')])
for i in range(1,6):
 d.add_page_break(); head(d,f'Pillar {i}')
 table(d,[('Please tell us about your pillar','Your response'),('1. What is your pillar about?',''),('2. What short description should appear at the top of the page?',''),('3. What important content should we highlight?\nExamples: programmes, achievements, advisors, partners, projects, resources',''),('4. Are there any sub-pillars or sections you want included?',''),('5. How would you like the page to be laid out?\nYou may attach a sketch, slide, or example website.',''),('6. Are there any specific features you want?\nExamples: event sign-up, contact form, image gallery, resource downloads, news, FAQs, interactive elements',''),('7. Are there any events you want included on the HealthX timeline?\nPlease include the event name, date, time, and link if available.',''),('8. Is there anything else we should know?','')])
 d.add_paragraph('Optional: attach any images, mock-ups, examples, or existing write-ups that would help us understand your preference.')
d.core_properties.title='Website Pillar Needs Form'; d.core_properties.subject='Simple initial requirements form for HealthX pillar teams'; d.save(OUT); print(OUT)
